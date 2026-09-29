from __future__ import annotations
import io, re, zipfile
from dataclasses import dataclass
from docx import Document
from pypdf import PdfReader
from ..domain.models import Segment

FILE_MAX_BYTES=5_242_880; PDF_MAX_PAGES=20; DOC_MAX_CHARS=60_000; DOCX_MAX_EXPANDED_BYTES=25_000_000; DOCX_MAX_PARAGRAPHS=5_000
class DocumentError(ValueError): pass
@dataclass
class ParsedDocument: segments:list[Segment]; page_count:int|None; warnings:list[str]

def _validate_size(data:bytes):
    if len(data)>FILE_MAX_BYTES: raise DocumentError("File exceeds the 5 MiB limit")
def _finalize(segments:list[Segment],pages:int|None,warnings:list[str]):
    chars=sum(len(s.text) for s in segments)
    if chars>DOC_MAX_CHARS: raise DocumentError("Extracted text exceeds 60,000 characters; the file was not truncated")
    if chars<200: raise DocumentError("Extraction is unreliable or the document appears scanned; upload a digital text file")
    return ParsedDocument(segments,pages,warnings)
def parse_pdf(data:bytes)->ParsedDocument:
    _validate_size(data)
    if not data.startswith(b"%PDF-"): raise DocumentError("File signature is not a PDF")
    try: reader=PdfReader(io.BytesIO(data),strict=True)
    except Exception as exc: raise DocumentError("Malformed PDF") from exc
    if reader.is_encrypted: raise DocumentError("Encrypted PDFs are not supported")
    if len(reader.pages)>PDF_MAX_PAGES: raise DocumentError("PDF exceeds the 20 page limit")
    segments=[]; sparse=0
    for index,page in enumerate(reader.pages,1):
        text=_normalize(page.extract_text() or "")
        if len(re.sub(r"\s","",text))<30: sparse+=1
        if text: segments.append(Segment(segment_id=f"pdf-p{index}",source_type="pdf",page=index,text=text))
    if reader.pages and sparse/len(reader.pages)>.2: raise DocumentError("Too many pages lack readable text; scans and mixed image PDFs are unsupported")
    return _finalize(segments,len(reader.pages),[])
def parse_docx(data:bytes)->ParsedDocument:
    _validate_size(data)
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            expanded=sum(i.file_size for i in archive.infolist())
            if expanded>DOCX_MAX_EXPANDED_BYTES: raise DocumentError("DOCX expanded content exceeds the 25 MB safety limit")
        doc=Document(io.BytesIO(data))
    except DocumentError: raise
    except Exception as exc: raise DocumentError("Malformed or unsupported DOCX") from exc
    if len(doc.paragraphs)>DOCX_MAX_PARAGRAPHS: raise DocumentError("DOCX exceeds the 5,000 paragraph limit")
    segments=[]
    for i,p in enumerate(doc.paragraphs):
        text=_normalize(p.text)
        if text: segments.append(Segment(segment_id=f"docx-p{i}",source_type="docx",paragraph_index=i,text=text))
    for ti,table in enumerate(doc.tables):
        for ri,row in enumerate(table.rows):
            for ci,cell in enumerate(row.cells):
                text=_normalize(cell.text)
                if text: segments.append(Segment(segment_id=f"docx-t{ti}-r{ri}-c{ci}",source_type="docx",table_cell=f"{ti}:{ri}:{ci}",text=text))
    return _finalize(segments,None,[])
def _normalize(text:str)->str: return re.sub(r"[ \t]+"," ",re.sub(r"\r\n?","\n",text)).strip()

def redact_segments(segments:list[Segment],candidate_name:str|None=None)->tuple[list[Segment],list[str]]:
    patterns=[(re.compile(r"\b[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}\b"),"[EMAIL]"),(re.compile(r"(?<!\d)(?:\+?\d[\d ()-]{7,}\d)(?!\d)"),"[PHONE]")]
    if candidate_name: patterns.append((re.compile(re.escape(candidate_name),re.I),"[NAME]"))
    out=[]; warnings=[]
    for segment in segments:
        text=segment.text
        for pattern,replacement in patterns: text=pattern.sub(replacement,text)
        if segment.paragraph_index==0 and "[" not in text: warnings.append(f"Review possible identity in {segment.segment_id}")
        out.append(segment.model_copy(update={"text":text}))
    return out,warnings
